"""政务云资产 API：按资源类型分表，Excel「原始数据」列一比一入库。

不导入 Sheet「数据统计」及统计类表头；分布统计由库内字段自行聚合（前端暂不展示）。
"""
from __future__ import annotations

import io
import json
import logging
import time
from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from sqlalchemy import func, or_
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from app.core.audit import get_client_ip, log_audit
from app.core.govcloud_schema import (
    FIELD_LABELS,
    RESERVED_COLUMNS,
    RESOURCE_EXCEL_FIELDS,
    SOURCE_ID_KEYS,
    excel_attr,
    skip_excel_header,
)
from app.core.timezone import beijing_now
from app.database import get_db
from app.dependencies import require_permission
from app.models.govcloud_asset import GovCloudImportBatch, RESOURCE_MODELS
from app.models.user import User

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/govcloud", tags=["govcloud"])

RESOURCE_TYPES = {
    "cloud_host": {
        "label": "云主机",
        "excel_name": "云主机资产",
        "source_group": "itsm_ops",
        "filename_hints": ("云主机", "cloud_host"),
        "search_fields": (
            "instance_name", "inst_name", "ip_address", "floatip_address",
            "user_department", "applicant_department", "source_id",
        ),
    },
    "bare_metal": {
        "label": "裸金属",
        "excel_name": "裸金属资产",
        "source_group": "itsm_ops",
        "filename_hints": ("裸金属", "bare_metal"),
        "search_fields": (
            "instance_name", "inst_name", "ip_address", "floatip_address",
            "user_department", "applicant_department", "source_id",
        ),
    },
    "e_government_network": {
        "label": "网络资源",
        "excel_name": "网络资源资产",
        "source_group": "itsm_ops",
        "filename_hints": ("网络资源", "e_government_network", "network"),
        "search_fields": (
            "instance_name", "inst_name", "name", "ip_address",
            "user_department", "applicant_department", "source_id",
        ),
    },
    "elastic_ip": {
        "label": "弹性IP",
        "excel_name": "弹性IP资产",
        "source_group": "itsm_ops",
        "filename_hints": ("弹性IP", "弹性ip", "elastic_ip"),
        "search_fields": (
            "instance_name", "inst_name", "ip_address", "floatip_address",
            "user_department", "applicant_department", "source_id",
        ),
    },
}

SOURCE_GROUPS = {
    "itsm_ops": {
        "label": "综合运管平台资产发现",
        "index": "",
        "enabled": True,
        "sort_order": 1,
        "description": "综合运管平台每周全量爬取，对应云主机 / 裸金属 / 网络资源 / 弹性IP 四张 Excel",
    },
    "intranet_mapping": {
        "label": "内网资产测绘",
        "index": "1.1",
        "enabled": False,
        "sort_order": 2,
        "description": "内网资产测绘结果，后续接入",
    },
    "qingteng_host": {
        "label": "青藤主机资产发现",
        "index": "1.2",
        "enabled": False,
        "sort_order": 3,
        "description": "青藤主机资产发现结果，后续接入",
    },
    "jiaotu_yunsuo": {
        "label": "椒图云锁主机资产发现",
        "index": "1.3",
        "enabled": False,
        "sort_order": 4,
        "description": "椒图云锁主机资产发现结果，后续接入",
    },
    "tianqing_endpoint": {
        "label": "天擎终端资产发现",
        "index": "1.4",
        "enabled": False,
        "sort_order": 5,
        "description": "天擎终端资产发现结果，后续接入",
    },
}

MAX_ROWS = 250000
MAX_FILE_SIZE = 200 * 1024 * 1024
UPSERT_CHUNK = 500


def _require_type(resource_type: str) -> dict:
    meta = RESOURCE_TYPES.get(resource_type)
    if not meta:
        raise HTTPException(status_code=400, detail="不支持的资源类型")
    return meta


def _model(resource_type: str):
    _require_type(resource_type)
    return RESOURCE_MODELS[resource_type]


def _cell_value(val) -> str:
    if val is None:
        return ""
    if isinstance(val, (dict, list)):
        return json.dumps(val, ensure_ascii=False)
    if isinstance(val, bool):
        return "true" if val else "false"
    return str(val).strip()


def _parse_excel(raw: bytes) -> list[dict]:
    """只读 Sheet「原始数据」，跳过「数据统计」页和统计类表头。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = None
    for sheet in wb.worksheets:
        title = (sheet.title or "").strip()
        if title == "数据统计":
            continue
        if title == "原始数据":
            ws = sheet
            break
    if ws is None:
        for sheet in wb.worksheets:
            if (sheet.title or "").strip() != "数据统计":
                ws = sheet
                break
    if ws is None:
        return []
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [_cell_value(h) for h in next(rows_iter)]
    except StopIteration:
        return []
    result: list[dict] = []
    for row in rows_iter:
        item = {}
        empty = True
        for i, key in enumerate(headers):
            if skip_excel_header(key):
                continue
            val = _cell_value(row[i] if i < len(row) else "")
            if val:
                empty = False
            item[key] = val
        if not empty:
            result.append(item)
        if len(result) > MAX_ROWS:
            raise HTTPException(status_code=413, detail=f"行数过多，最多 {MAX_ROWS} 行")
    return result


def _source_id(row: dict) -> str:
    for key in SOURCE_ID_KEYS:
        val = _cell_value(row.get(key))
        if val:
            return val[:128]
    ip = _cell_value(row.get("ip_address") or row.get("ip"))
    name = _cell_value(row.get("instance_name") or row.get("inst_name") or row.get("name"))
    if ip or name:
        return f"{ip}|{name}"[:128]
    return ""


def _row_to_values(resource_type: str, row: dict, batch_id: int, now) -> dict:
    fields = RESOURCE_EXCEL_FIELDS[resource_type]
    known = set(fields)
    values = {
        "source_id": _source_id(row),
        "batch_id": batch_id,
        "created_at": now,
        "updated_at": now,
    }
    extra = {}
    for key, raw in row.items():
        if skip_excel_header(key):
            continue
        val = _cell_value(raw)
        if key in RESERVED_COLUMNS:
            extra[key] = val
            continue
        if key in known:
            values[excel_attr(key)] = val
        else:
            extra[key] = val
    for field in fields:
        attr = excel_attr(field)
        values.setdefault(attr, "")
    values["extra_excel"] = extra
    return values


def infer_resource_type(filename: str, explicit: str) -> str:
    if explicit:
        _require_type(explicit)
        return explicit
    name = (filename or "").lower()
    for code, meta in RESOURCE_TYPES.items():
        for hint in meta["filename_hints"]:
            if hint.lower() in name:
                return code
    raise HTTPException(
        status_code=400,
        detail="无法从文件名识别资源类型，请选择：云主机 / 裸金属 / 网络资源 / 弹性IP",
    )


def _excel_columns(resource_type: str) -> list[dict]:
    cols = []
    for key in RESOURCE_EXCEL_FIELDS[resource_type]:
        cols.append({"key": key, "label": FIELD_LABELS.get(key, key)})
    return cols


def _type_payload(code: str, meta: dict) -> dict:
    return {
        "code": code,
        "label": meta["label"],
        "excel_name": meta.get("excel_name") or meta["label"],
        "source_group": meta.get("source_group", "itsm_ops"),
        "excel_columns": _excel_columns(code),
    }


@router.get("/meta")
def govcloud_meta(
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    types = [_type_payload(code, meta) for code, meta in RESOURCE_TYPES.items()]
    groups = []
    for gcode, gmeta in SOURCE_GROUPS.items():
        groups.append({
            "code": gcode,
            "label": gmeta["label"],
            "index": gmeta.get("index", ""),
            "enabled": bool(gmeta.get("enabled", True)),
            "sort_order": gmeta.get("sort_order", 0),
            "description": gmeta.get("description", ""),
            "tables": [t for t in types if t["source_group"] == gcode],
        })
    groups.sort(key=lambda g: g["sort_order"])
    return {"types": types, "groups": groups, "field_labels": FIELD_LABELS}


@router.get("/summary")
def govcloud_summary(
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    groups = []
    for gcode, gmeta in SOURCE_GROUPS.items():
        tables = []
        for code, meta in RESOURCE_TYPES.items():
            if meta.get("source_group") != gcode:
                continue
            model = RESOURCE_MODELS[code]
            live = int(db.query(func.count(model.id)).scalar() or 0)
            last_batch = (
                db.query(GovCloudImportBatch)
                .filter(GovCloudImportBatch.resource_type == code)
                .order_by(GovCloudImportBatch.id.desc())
                .first()
            )
            unique_ids = last_batch.unique_source_ids if last_batch else 0
            consistent = (live == unique_ids and (last_batch.skipped == 0 if last_batch else True))
            if last_batch is None:
                consistent = live == 0
            tables.append({
                "code": code,
                "label": meta["label"],
                "excel_name": meta.get("excel_name") or meta["label"],
                "excel_columns": _excel_columns(code),
                "total": live,
                "consistent": consistent,
                "last_batch": last_batch.to_dict() if last_batch else None,
            })
        groups.append({
            "code": gcode,
            "label": gmeta["label"],
            "index": gmeta.get("index", ""),
            "enabled": bool(gmeta.get("enabled", True)),
            "sort_order": gmeta.get("sort_order", 0),
            "description": gmeta.get("description", ""),
            "tables": tables,
        })
    groups.sort(key=lambda g: g["sort_order"])
    return {"groups": groups}


@router.get("/assets")
def list_govcloud_assets(
    resource_type: str = Query(..., description="资源类型"),
    keyword: str = Query("", description="名称/IP/部门/source_id 模糊搜索"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    meta = _require_type(resource_type)
    model = _model(resource_type)
    query = db.query(model)
    kw = keyword.strip()
    if kw:
        like = f"%{kw}%"
        clauses = [model.source_id.ilike(like)]
        if hasattr(model, "cmdb_id"):
            clauses.append(model.cmdb_id.ilike(like))
        for field in meta.get("search_fields") or []:
            if field in ("source_id", "_id"):
                continue
            attr = excel_attr(field)
            col = getattr(model, attr, None)
            if col is not None:
                clauses.append(col.ilike(like))
        query = query.filter(or_(*clauses))

    total = query.count()
    rows = (
        query.order_by(model.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "columns": _excel_columns(resource_type),
        "items": [r.to_dict() for r in rows],
    }


@router.get("/assets/{resource_type}/{asset_id}")
def get_govcloud_asset(
    resource_type: str,
    asset_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    model = _model(resource_type)
    rec = db.query(model).filter(model.id == asset_id).first()
    if rec is None:
        raise HTTPException(status_code=404, detail="记录不存在")
    return rec.to_dict()


@router.get("/batches")
def list_batches(
    resource_type: str = Query("", description="可选，按类型筛选"),
    limit: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    query = db.query(GovCloudImportBatch)
    if resource_type:
        _require_type(resource_type)
        query = query.filter(GovCloudImportBatch.resource_type == resource_type)
    rows = query.order_by(GovCloudImportBatch.id.desc()).limit(limit).all()
    return {"items": [r.to_dict() for r in rows]}


@router.get("/consistency")
def check_consistency(
    resource_type: str = Query(..., description="资源类型"),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "view")),  # noqa: ARG001
) -> dict:
    _require_type(resource_type)
    model = _model(resource_type)
    live = int(db.query(func.count(model.id)).scalar() or 0)
    last_full = (
        db.query(GovCloudImportBatch)
        .filter(
            GovCloudImportBatch.resource_type == resource_type,
            or_(
                GovCloudImportBatch.import_mode == "full",
                GovCloudImportBatch.import_mode.is_(None),
                GovCloudImportBatch.import_mode == "",
            ),
        )
        .order_by(GovCloudImportBatch.id.desc())
        .first()
    )
    last_batch = (
        db.query(GovCloudImportBatch)
        .filter(GovCloudImportBatch.resource_type == resource_type)
        .order_by(GovCloudImportBatch.id.desc())
        .first()
    )
    if last_full is None:
        return {
            "resource_type": resource_type,
            "live_count": live,
            "excel_rows": last_batch.excel_rows if last_batch else 0,
            "unique_source_ids": last_batch.unique_source_ids if last_batch else 0,
            "consistent": live == 0,
            "last_batch": last_batch.to_dict() if last_batch else None,
            "message": (
                "尚未做过全量导入，无法按全量爬取校验一致性"
                if last_batch
                else "尚未导入过该类型，库内应为空才算一致"
            ),
        }
    consistent = live == last_full.unique_source_ids and last_full.skipped == 0
    if last_full.consistent != int(consistent) or last_full.live_count != live:
        last_full.consistent = int(consistent)
        last_full.live_count = live
        db.commit()
    return {
        "resource_type": resource_type,
        "live_count": live,
        "excel_rows": last_full.excel_rows,
        "unique_source_ids": last_full.unique_source_ids,
        "skipped": last_full.skipped,
        "consistent": consistent,
        "last_batch": last_batch.to_dict() if last_batch else None,
        "last_full": last_full.to_dict(),
        "message": (
            "与最近一次全量导入一致"
            if consistent
            else f"与最近一次全量不一致：库内 {live} 条，全量唯一 ID {last_full.unique_source_ids} 条"
        ),
    }


@router.post("/import")
async def import_govcloud_excel(
    file: UploadFile = File(..., description="爬虫产出的 xlsx"),
    resource_type: str = Query("", description="cloud_host / bare_metal / e_government_network / elastic_ip"),
    mode: str = Query("full", description="导入模式：full 全量对齐 / incremental 增量 upsert 不删除"),
    request: Request = None,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("asset_discovery", "edit")),
) -> dict:
    filename = file.filename or "upload.xlsx"
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx（爬虫输出格式）")
    import_mode = mode if mode in ("full", "incremental") else "full"
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="文件为空")
    if len(raw) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="文件超过 200MB")

    rtype = infer_resource_type(filename, resource_type)
    model = _model(rtype)
    started = time.perf_counter()
    rows = _parse_excel(raw)
    if not rows:
        raise HTTPException(status_code=400, detail="未解析到数据行（请确认 Sheet「原始数据」，统计页不会导入）")

    skipped = 0
    seen: dict[str, dict] = {}
    duplicate = 0
    for row in rows:
        sid = _source_id(row)
        if not sid:
            skipped += 1
            continue
        if sid in seen:
            duplicate += 1
        seen[sid] = row

    unique_rows = list(seen.values())
    unique_ids = set(seen.keys())
    existing_ids = {r[0] for r in db.query(model.source_id).all()}
    inserted = len(unique_ids - existing_ids)
    updated = len(unique_ids & existing_ids)
    deleted = len(existing_ids - unique_ids) if import_mode == "full" else 0

    batch = GovCloudImportBatch(
        resource_type=rtype,
        filename=filename[:300],
        file_size=len(raw),
        excel_rows=len(rows),
        unique_source_ids=len(unique_ids),
        inserted=inserted,
        updated=updated,
        deleted=deleted,
        skipped=skipped,
        live_count=0,
        consistent=0,
        import_mode=import_mode,
        created_by=user.id,
        created_at=beijing_now(),
    )
    db.add(batch)
    db.flush()

    now = beijing_now()
    values = [_row_to_values(rtype, row, batch.id, now) for row in unique_rows]
    excel_attrs = [excel_attr(f) for f in RESOURCE_EXCEL_FIELDS[rtype]]
    update_cols = [*excel_attrs, "extra_excel", "batch_id", "updated_at"]

    for i in range(0, len(values), UPSERT_CHUNK):
        chunk = values[i:i + UPSERT_CHUNK]
        stmt = insert(model).values(chunk)
        excluded = stmt.excluded
        stmt = stmt.on_conflict_do_update(
            index_elements=["source_id"],
            set_={col: getattr(excluded, col) for col in update_cols},
        )
        db.execute(stmt)

    if import_mode == "full":
        db.query(model).filter(model.batch_id != batch.id).delete(synchronize_session=False)

    live = int(db.query(func.count(model.id)).scalar() or 0)
    consistent = live == len(unique_ids) and skipped == 0 if import_mode == "full" else skipped == 0
    notes = []
    notes.append("全量导入：已删除文件中不存在的旧记录" if import_mode == "full" else "增量导入：仅新增/更新，未删除库内既有记录")
    notes.append("未导入 Sheet「数据统计」")
    if duplicate:
        notes.append(f"Excel 内重复 source_id {duplicate} 条，已按最后一条保留")
    if skipped:
        notes.append(f"缺少 _id/instance_id 跳过 {skipped} 条")
    if not consistent:
        notes.append(f"同步后库内 {live} 与唯一 ID {len(unique_ids)} 不一致")
    batch.live_count = live
    batch.consistent = int(consistent)
    batch.duration_ms = int((time.perf_counter() - started) * 1000)
    batch.note = "；".join(notes)
    db.commit()

    try:
        log_audit(
            db,
            user_id=user.id,
            username=user.username,
            action="import",
            resource_type="asset_discovery",
            resource_id=str(batch.id),
            detail=f"{rtype} {import_mode} excel={len(rows)} live={live}",
            ip_address=get_client_ip(request) if request else None,
            result="success",
        )
    except Exception as extra_exc:  # noqa: BLE001
        logger.warning("导入审计写入失败（忽略）: %s", extra_exc)
    logger.info(
        "政务云导入: mode=%s type=%s excel=%d unique=%d live=%d inserted=%d updated=%d deleted=%d consistent=%s",
        import_mode, rtype, len(rows), len(unique_ids), live, inserted, updated, deleted, consistent,
    )
    return {
        "resource_type": rtype,
        "import_mode": import_mode,
        "excel_rows": len(rows),
        "unique_source_ids": len(unique_ids),
        "inserted": inserted,
        "updated": updated,
        "deleted": deleted,
        "skipped": skipped,
        "live_count": live,
        "consistent": consistent,
        "duration_ms": batch.duration_ms,
        "note": batch.note,
        "batch": batch.to_dict(),
    }
